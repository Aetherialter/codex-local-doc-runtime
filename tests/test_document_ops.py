from pathlib import Path

import fitz
import openpyxl
import pytest
from docx import Document

from docrt.docx_ops import inspect_docx
from docrt.models import ErrorCode
from docrt.paths import ValidationError
from docrt.pdf_ops import inspect_pdf, render_pdf
from docrt.read_ops import read_pdf
from docrt.xlsx_ops import inspect_xlsx


def test_inspect_docx(tmp_path: Path):
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("hello docx")
    document.save(path)

    result = inspect_docx(path)

    assert result["paragraph_count"] == 1
    assert result["paragraphs"] == ["hello docx"]


def test_inspect_xlsx(tmp_path: Path):
    path = tmp_path / "sample.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "hello"
    workbook.save(path)

    result = inspect_xlsx(path)

    assert result["sheet_count"] == 1
    assert result["sheets"][0]["preview"][0][0] == "hello"


def test_inspect_and_render_pdf(tmp_path: Path):
    path = tmp_path / "sample.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "hello pdf")
    document.save(path)
    document.close()

    info = inspect_pdf(path)
    rendered = render_pdf(path, tmp_path / "rendered")

    assert info["page_count"] == 1
    assert info["has_text_layer"] is True
    assert info["needs_ocr"] is False
    assert len(rendered["pages"]) == 1


def test_pdf_without_text_layer_reports_ocr_needed(tmp_path: Path):
    path = tmp_path / "image-only.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()

    info = inspect_pdf(path)
    content = read_pdf(path)

    assert info["has_text_layer"] is False
    assert info["needs_ocr"] is True
    assert info["ocr_supported"] is False
    assert content["metadata"]["needs_ocr"] is True
    assert "OCR is not supported in v1.1" in content["warnings"][0]


def test_encrypted_pdf_is_rejected(tmp_path: Path):
    path = tmp_path / "encrypted.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "secret")
    document.save(
        path,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="user",
    )
    document.close()

    with pytest.raises(ValidationError) as exc_info:
        inspect_pdf(path)

    assert exc_info.value.error_code == ErrorCode.ENCRYPTED_FILE_UNSUPPORTED
